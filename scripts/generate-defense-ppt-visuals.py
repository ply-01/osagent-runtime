from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BLUE = RGBColor(0x0B, 0x6F, 0xB8)
DARK_BLUE = RGBColor(0x13, 0x3D, 0x61)
RED = RGBColor(0xDF, 0x32, 0x2B)
GREEN = RGBColor(0x17, 0x91, 0x76)
TEXT = RGBColor(0x1F, 0x2D, 0x3D)
MUTED = RGBColor(0x67, 0x78, 0x88)
LIGHT_BLUE = RGBColor(0xE9, 0xF3, 0xFB)
LIGHT_GREEN = RGBColor(0xE9, 0xF7, 0xF3)
LIGHT_RED = RGBColor(0xFE, 0xF0, 0xEF)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"


def _remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def clear_body(slide) -> None:
    title_element = slide.shapes.title._element if slide.shapes.title else None
    for shape in list(slide.shapes):
        if shape._element is title_element:
            continue
        if shape.is_placeholder:
            placeholder_type = shape.placeholder_format.type
            if placeholder_type == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
        if (
            shape.shape_type == MSO_SHAPE_TYPE.LINE
            and shape.top <= Inches(0.7)
            and shape.width >= Inches(12.0)
        ):
            continue
        if shape.top >= Inches(7.25) and shape.width >= Inches(12.0):
            continue
        _remove_shape(shape)


def set_title(slide, text: str) -> None:
    title = slide.shapes.title
    if title is None:
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.12), Inches(8.8), Inches(0.55))
    title.text = text
    paragraph = title.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.runs[0]
    run.font.name = FONT
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = BLUE


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 16,
    color: RGBColor = TEXT,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    margin: float = 0.04,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_round_rect(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = LIGHT_BLUE,
    radius_shape: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE,
    line_width: float = 1.2,
):
    shape = slide.shapes.add_shape(
        radius_shape,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_subtitle(slide, text: str) -> None:
    add_text(
        slide,
        text,
        1.1,
        0.78,
        11.1,
        0.44,
        size=18,
        color=TEXT,
        align=PP_ALIGN.CENTER,
    )


def add_official_frame(slide, page_number: int) -> None:
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0.615),
        Inches(13.333),
        Inches(0.015),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.35),
        Inches(13.333),
        Inches(0.15),
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = BLUE
    footer.line.fill.background()
    add_text(
        slide,
        str(page_number),
        12.06,
        6.92,
        0.28,
        0.3,
        size=12,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_pain_card(
    slide,
    number: str,
    heading: str,
    observation: str,
    impact: str,
    x: float,
    y: float,
) -> None:
    add_round_rect(slide, x, y, 5.82, 1.72, fill=WHITE, line=RGBColor(0xB9, 0xD4, 0xE8))
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.08),
        Inches(1.72),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x + 0.28),
        Inches(y + 0.26),
        Inches(0.54),
        Inches(0.54),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = LIGHT_RED
    circle.line.fill.background()
    add_text(slide, number, x + 0.28, y + 0.26, 0.54, 0.54, size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, heading, x + 1.0, y + 0.18, 4.45, 0.42, size=20, color=DARK_BLUE, bold=True)
    add_text(slide, f"现象  {observation}", x + 1.0, y + 0.65, 4.45, 0.42, size=13.5, color=TEXT)
    add_text(slide, f"影响  {impact}", x + 1.0, y + 1.08, 4.45, 0.42, size=13.5, color=RED)


def build_background_slide(slide) -> None:
    clear_body(slide)
    set_title(slide, "行业背景与核心痛点")
    add_subtitle(slide, "当 Agent 从单轮问答走向长期协作，记忆成为决定可用性的基础能力")
    add_pain_card(
        slide,
        "01",
        "记忆连续性不足",
        "会话结束后，偏好、项目状态和执行经验难以自然延续",
        "用户反复说明，长期任务难以接续",
        0.68,
        1.42,
    )
    add_pain_card(
        slide,
        "02",
        "记忆可信性不足",
        "临时要求、否定、推测与工具噪声被混合保存",
        "错误信息进入长期记忆并影响决策",
        6.83,
        1.42,
    )
    add_pain_card(
        slide,
        "03",
        "动态事实一致性不足",
        "事实随实体、项目、作用域和时间持续变化",
        "新旧事实冲突，缺少可解释的版本依据",
        0.68,
        3.42,
    )
    add_pain_card(
        slide,
        "04",
        "国产端侧治理不足",
        "通用云端组件难兼顾系统 SDK、本地数据、权限与审计",
        "原型能够演示，却难以在银河麒麟稳定落地",
        6.83,
        3.42,
    )
    add_round_rect(slide, 1.45, 5.58, 10.43, 0.58, fill=LIGHT_BLUE, line=RGBColor(0xB9, 0xD4, 0xE8))
    add_text(
        slide,
        "需要的不是更长的聊天记录，而是一套可形成、可更新、可检索、可遗忘、可审计的长期记忆。",
        1.62,
        5.64,
        10.1,
        0.43,
        size=16,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_metric_card(
    slide,
    title: str,
    before: str,
    current: str,
    delta: str,
    x: float,
    y: float,
) -> None:
    add_round_rect(slide, x, y, 3.55, 2.02, fill=WHITE, line=RGBColor(0xB9, 0xD4, 0xE8))
    add_text(slide, title, x + 0.24, y + 0.16, 2.08, 0.38, size=16, color=DARK_BLUE, bold=True)
    add_round_rect(slide, x + 2.43, y + 0.16, 0.88, 0.36, fill=LIGHT_GREEN, line=LIGHT_GREEN)
    add_text(slide, delta, x + 2.45, y + 0.17, 0.84, 0.32, size=10.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, before, x + 0.18, y + 0.68, 1.18, 0.58, size=22, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", x + 1.34, y + 0.68, 0.42, 0.58, size=24, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, current, x + 1.72, y + 0.58, 1.58, 0.7, size=27, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "优化前", x + 0.32, y + 1.38, 0.9, 0.32, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "当前保留版本", x + 1.76, y + 1.38, 1.46, 0.32, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)


def add_status_card(slide, value: str, label: str, x: float, y: float, *, green: bool = False) -> None:
    fill = LIGHT_GREEN if green else LIGHT_BLUE
    color = GREEN if green else BLUE
    add_round_rect(slide, x, y, 3.78, 0.98, fill=fill, line=fill)
    add_text(slide, value, x + 0.18, y + 0.12, 1.42, 0.56, size=25, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 1.62, y + 0.16, 1.92, 0.5, size=13.5, color=TEXT, bold=True)


def build_results_slide(slide) -> None:
    clear_body(slide)
    set_title(slide, "验证结果：记忆质量与原生工程能力同步提升")
    add_subtitle(slide, "冻结数据 · 标签隔离 · 同协议 A/B · 退化候选回退")
    add_text(slide, "记忆质量  |  KylinMemoryEval-v5", 0.68, 1.28, 7.55, 0.34, size=15, color=DARK_BLUE, bold=True)
    add_metric_card(slide, "偏好提取 Accuracy", "90.0%", "96.6%", "+6.6pp", 0.68, 1.67)
    add_metric_card(slide, "Value grounding", "85.6%", "93.2%", "+7.6pp", 4.48, 1.67)
    add_text(slide, "工程闭环  |  银河麒麟 V11", 8.55, 1.28, 4.1, 0.34, size=15, color=DARK_BLUE, bold=True)
    add_status_card(slide, "440ms", "10k 在线检索 p95", 8.55, 1.67)
    add_status_card(slide, "7 / 7", "系统严格预检", 8.55, 2.83, green=True)
    add_status_card(slide, "CRUD", "五项原生操作通过", 8.55, 3.99, green=True)

    add_round_rect(slide, 0.68, 4.12, 7.35, 1.14, fill=LIGHT_GRAY, line=RGBColor(0xDD, 0xE4, 0xEA))
    add_text(slide, "索引效率  |  PERMA 增量语义索引 p95", 0.94, 4.28, 3.35, 0.36, size=14, color=DARK_BLUE, bold=True)
    add_text(slide, "7.07s", 4.24, 4.2, 1.08, 0.54, size=22, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 5.3, 4.2, 0.42, 0.54, size=24, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "3.32s", 5.72, 4.13, 1.3, 0.62, size=27, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 7.02, 4.28, 0.84, 0.4, fill=LIGHT_GREEN, line=LIGHT_GREEN)
    add_text(slide, "-53.0%", 7.04, 4.3, 0.8, 0.35, size=10.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 0.9, 5.62, 11.55, 0.5, fill=WHITE, line=RGBColor(0xD4, 0xE2, 0xED))
    add_text(
        slide,
        "数据口径：质量指标来自 v5 同协议 A/B；索引指标来自 PERMA 同轮实验；在线性能来自 10,000 条 Core facts、500 次查询。",
        1.03,
        5.67,
        11.28,
        0.36,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def build_official_results_slide(slide) -> None:
    clear_body(slide)
    set_title(slide, "量化验证：同协议 A/B 下的记忆质量提升")
    add_subtitle(slide, "KylinMemoryEval-v5 冻结复用回归 · 预测完成后加载标签 · 退化候选回退")

    add_round_rect(slide, 0.68, 1.4, 8.18, 4.15, fill=WHITE, line=RGBColor(0xC8, 0xD9, 0xE6))
    add_text(slide, "三项质量指标（%）", 0.96, 1.56, 2.0, 0.34, size=15, color=DARK_BLUE, bold=True)
    add_text(slide, "■ 优化前", 6.35, 1.58, 0.92, 0.28, size=10, color=MUTED)
    add_text(slide, "■ 当前", 7.35, 1.58, 0.84, 0.28, size=10, color=BLUE)

    chart_x = 3.0
    chart_width = 4.55
    chart_top = 2.12
    chart_bottom = 5.08
    for tick in (0, 25, 50, 75, 100):
        x = chart_x + chart_width * tick / 100
        grid = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(chart_top),
            Inches(0.008),
            Inches(chart_bottom - chart_top),
        )
        grid.fill.solid()
        grid.fill.fore_color.rgb = RGBColor(0xE4, 0xE9, 0xEE)
        grid.line.fill.background()
        add_text(slide, str(tick), x - 0.18, 5.08, 0.36, 0.24, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)

    comparisons = (
        ("偏好提取", 90.0, 96.6, "+6.6pp", 2.25),
        ("Value grounding", 85.6, 93.2, "+7.6pp", 3.18),
        ("规范化属性", 61.5, 91.0, "+29.5pp", 4.11),
    )
    for label, before, current, delta, y in comparisons:
        add_text(slide, label, 0.96, y + 0.08, 1.72, 0.38, size=13.5, color=DARK_BLUE, bold=True)
        before_width = chart_width * before / 100
        current_width = chart_width * current / 100
        add_round_rect(
            slide,
            chart_x,
            y,
            before_width,
            0.18,
            fill=RGBColor(0xD7, 0xDF, 0xE7),
            line=RGBColor(0xD7, 0xDF, 0xE7),
            radius_shape=MSO_SHAPE.RECTANGLE,
        )
        add_round_rect(
            slide,
            chart_x,
            y + 0.25,
            current_width,
            0.28,
            fill=BLUE,
            line=BLUE,
            radius_shape=MSO_SHAPE.RECTANGLE,
        )
        add_text(slide, f"{before:.1f}", chart_x + before_width - 0.62, y - 0.02, 0.56, 0.22, size=8.5, color=TEXT, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, f"{current:.1f}", chart_x + current_width - 0.62, y + 0.27, 0.56, 0.22, size=9, color=WHITE, bold=True, align=PP_ALIGN.RIGHT)
        add_round_rect(slide, 7.72, y + 0.06, 0.9, 0.38, fill=LIGHT_GREEN, line=LIGHT_GREEN)
        add_text(slide, delta, 7.75, y + 0.08, 0.84, 0.32, size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 9.08, 1.4, 3.55, 4.15, fill=WHITE, line=RGBColor(0xC8, 0xD9, 0xE6))
    add_text(slide, "10k 在线检索 p95", 9.4, 1.65, 2.9, 0.38, size=15, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "440.464", 9.48, 2.14, 2.46, 0.62, size=29, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "ms", 11.76, 2.32, 0.42, 0.28, size=12, color=MUTED, bold=True)

    gauge_x = 9.52
    gauge_y = 3.08
    gauge_width = 2.66
    add_round_rect(slide, gauge_x, gauge_y, gauge_width, 0.22, fill=RGBColor(0xE2, 0xE8, 0xEE), line=RGBColor(0xE2, 0xE8, 0xEE), radius_shape=MSO_SHAPE.RECTANGLE)
    add_round_rect(slide, gauge_x, gauge_y, gauge_width * 440.464 / 600, 0.22, fill=BLUE, line=BLUE, radius_shape=MSO_SHAPE.RECTANGLE)
    threshold_x = gauge_x + gauge_width * 500 / 600
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(threshold_x), Inches(gauge_y - 0.11), Inches(0.025), Inches(0.44))
    marker.fill.solid()
    marker.fill.fore_color.rgb = GREEN
    marker.line.fill.background()
    add_text(slide, "0", gauge_x - 0.08, gauge_y + 0.3, 0.2, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "500", threshold_x - 0.24, gauge_y + 0.3, 0.48, 0.22, size=8.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "600ms", gauge_x + gauge_width - 0.44, gauge_y + 0.3, 0.58, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 9.72, 3.76, 2.25, 0.42, fill=LIGHT_GREEN, line=LIGHT_GREEN)
    add_text(slide, "位于 500ms 门槛内", 9.78, 3.79, 2.12, 0.34, size=11, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Windows CPU · 500 次查询", 9.42, 4.38, 2.88, 0.28, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "跨运行缓存波动不计作算法收益", 9.35, 4.75, 3.0, 0.28, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 1.02, 5.78, 11.25, 0.46, fill=WHITE, line=RGBColor(0xD4, 0xE2, 0xED))
    add_text(
        slide,
        "协议规模：500 条偏好样本 · 10,000 条检索语料 · 500 次检索 Query · 冻结数据仅作稳定复用回归",
        1.15,
        5.82,
        10.98,
        0.34,
        size=10.8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text
    )


def replace_text_preserving_runs(slide, old: str, new: str) -> int:
    replacements = 0
    pending = list(slide.shapes)
    while pending:
        shape = pending.pop()
        if hasattr(shape, "shapes"):
            pending.extend(shape.shapes)
        if not hasattr(shape, "text_frame") or not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old not in run.text:
                    continue
                run.text = run.text.replace(old, new)
                replacements += 1
    return replacements


def overlay_engineering_regression_card(slide) -> None:
    patch = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(6.94),
        Inches(2.39),
        Inches(2.56),
        Inches(1.02),
    )
    patch.fill.solid()
    patch.fill.fore_color.rgb = RGBColor(0xE4, 0xF2, 0xFA)
    patch.line.fill.background()
    add_text(slide, "346 passed", 7.08, 2.53, 2.18, 0.5, size=23, color=BLUE, bold=True)
    add_text(slide, "最近一次完整归档回归", 7.08, 3.04, 2.18, 0.26, size=9.5, color=DARK_BLUE, bold=True)


def move_slide(prs: Presentation, slide, index: int) -> None:
    slide_id_list = prs.slides._sldIdLst
    slide_id = next(
        item for item in slide_id_list if prs.part.related_part(item.rId) is slide.part
    )
    slide_id_list.remove(slide_id)
    slide_id_list.insert(index, slide_id)


def add_roadmap_stage(
    slide,
    number: str,
    horizon: str,
    title: str,
    items: list[str],
    x: float,
    *,
    fill: RGBColor,
    accent: RGBColor,
) -> None:
    add_round_rect(slide, x, 1.62, 3.65, 3.62, fill=WHITE, line=RGBColor(0xC8, 0xD9, 0xE6))
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.24), Inches(1.83), Inches(0.62), Inches(0.62))
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.color.rgb = accent
    add_text(slide, number, x + 0.24, 1.83, 0.62, 0.62, size=16, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, horizon, x + 1.0, 1.76, 2.3, 0.32, size=13, color=MUTED)
    add_text(slide, title, x + 1.0, 2.08, 2.3, 0.44, size=21, color=DARK_BLUE, bold=True)
    for index, item in enumerate(items):
        y = 2.78 + index * 0.62
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.42), Inches(y + 0.12), Inches(0.11), Inches(0.11))
        dot.fill.solid()
        dot.fill.fore_color.rgb = accent
        dot.line.fill.background()
        add_text(slide, item, x + 0.7, y, 2.58, 0.38, size=13.5, color=TEXT)


def build_outlook_slide(prs: Presentation, layout=None):
    slide = prs.slides.add_slide(layout or prs.slide_layouts[3])
    clear_body(slide)
    set_title(slide, "未来展望：从长期记忆底座走向国产端侧智能协作平台")
    add_subtitle(slide, "沿着记忆深化、产品协同和国产生态三条路径持续演进")
    add_roadmap_stage(
        slide,
        "01",
        "近期",
        "记忆能力深化",
        ["统一状态与证据表示", "跨语言分解与重排", "增量索引与记忆审计"],
        0.7,
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_text(slide, "→", 4.42, 3.0, 0.34, 0.5, size=24, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_roadmap_stage(
        slide,
        "02",
        "中期",
        "产品与多端协同",
        ["插件权限与隔离执行", "手机与电脑安全配对", "指令审批、回执与审计"],
        4.84,
        fill=LIGHT_GREEN,
        accent=GREEN,
    )
    add_text(slide, "→", 8.56, 3.0, 0.34, 0.5, size=24, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_roadmap_stage(
        slide,
        "03",
        "长期",
        "国产智能体生态",
        ["适配国产 CPU 与终端", "统一系统 AI 能力接口", "记忆服务赋能行业应用"],
        8.98,
        fill=LIGHT_RED,
        accent=RED,
    )
    add_round_rect(slide, 1.0, 5.65, 11.3, 0.56, fill=LIGHT_BLUE, line=RGBColor(0xB9, 0xD4, 0xE8))
    add_text(
        slide,
        "让每一台国产桌面终端都拥有可持续学习、可信执行、由用户掌控的个人智能体基础设施。",
        1.15,
        5.7,
        11.0,
        0.4,
        size=16,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    return slide


def build_official_mobile_outlook_slide(prs: Presentation, layout=None):
    slide = prs.slides.add_slide(layout or prs.slide_layouts[3])
    clear_body(slide)
    set_title(slide, "未来展望：面向个人设备的可信多端 OS Agent")
    add_subtitle(slide, "手机负责发起与审批，银河麒麟电脑负责本地执行，OSagent 连接身份、记忆与审计")

    # Mobile terminal
    add_round_rect(slide, 0.72, 1.42, 2.46, 3.88, fill=WHITE, line=BLUE, line_width=1.8)
    speaker = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.55), Inches(1.61), Inches(0.78), Inches(0.08))
    speaker.fill.solid()
    speaker.fill.fore_color.rgb = DARK_BLUE
    speaker.line.fill.background()
    add_round_rect(slide, 0.96, 1.85, 1.98, 2.82, fill=LIGHT_BLUE, line=RGBColor(0xC6, 0xDC, 0xEC))
    add_text(slide, "手机端", 1.18, 2.02, 1.54, 0.36, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "OSagent Mobile", 1.08, 2.35, 1.76, 0.45, size=17, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)
    mobile_items = (("01", "发起任务"), ("02", "高风险审批"), ("03", "进度与结果"))
    for index, (number, label) in enumerate(mobile_items):
        y = 3.0 + index * 0.48
        add_round_rect(slide, 1.16, y, 0.42, 0.3, fill=WHITE, line=WHITE)
        add_text(slide, number, 1.18, y + 0.01, 0.38, 0.26, size=9, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, 1.68, y, 1.05, 0.3, size=11.5, color=TEXT)
    home = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.82), Inches(4.85), Inches(0.24), Inches(0.24))
    home.fill.background()
    home.line.color.rgb = MUTED

    # Secure relay and device routing
    add_round_rect(slide, 4.18, 1.72, 4.0, 3.38, fill=WHITE, line=GREEN, line_width=1.8)
    shield = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.84), Inches(1.96), Inches(0.68), Inches(0.68))
    shield.fill.solid()
    shield.fill.fore_color.rgb = LIGHT_GREEN
    shield.line.color.rgb = GREEN
    add_text(slide, "✓", 5.84, 1.96, 0.68, 0.68, size=21, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "安全中继与设备路由", 4.7, 2.72, 2.96, 0.42, size=18, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)
    relay_items = (
        ("账户", "OIDC 身份体系"),
        ("配对", "二维码绑定个人电脑"),
        ("通道", "电脑主动建立出站 WSS"),
        ("授权", "tenant · user · device · command"),
    )
    for index, (label, value) in enumerate(relay_items):
        y = 3.32 + index * 0.4
        add_round_rect(slide, 4.55, y, 0.64, 0.28, fill=LIGHT_GREEN, line=LIGHT_GREEN)
        add_text(slide, label, 4.57, y + 0.01, 0.6, 0.24, size=9.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, value, 5.37, y - 0.01, 2.42, 0.3, size=10.5, color=TEXT)

    # Kylin PC terminal
    add_round_rect(slide, 9.08, 1.52, 3.55, 3.08, fill=WHITE, line=BLUE, line_width=1.8)
    add_text(slide, "银河麒麟 V11 电脑端", 9.4, 1.78, 2.9, 0.42, size=17, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)
    pc_items = (
        ("Agent Runtime", BLUE, LIGHT_BLUE),
        ("OSagent Memory", GREEN, LIGHT_GREEN),
        ("插件与系统工具", RED, LIGHT_RED),
    )
    for index, (label, accent, fill) in enumerate(pc_items):
        y = 2.5 + index * 0.57
        add_round_rect(slide, 9.52, y, 2.66, 0.4, fill=fill, line=fill)
        add_text(slide, label, 9.62, y + 0.03, 2.46, 0.32, size=11.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    stand = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.68), Inches(4.6), Inches(0.36), Inches(0.34))
    stand.fill.solid()
    stand.fill.fore_color.rgb = DARK_BLUE
    stand.line.fill.background()
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.24), Inches(4.91), Inches(1.24), Inches(0.07))
    base.fill.solid()
    base.fill.fore_color.rgb = DARK_BLUE
    base.line.fill.background()

    add_text(slide, "指令 / 审批", 3.22, 2.22, 0.92, 0.28, size=9.5, color=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 3.35, 2.48, 0.62, 0.5, size=25, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "进度 / 回执", 3.22, 3.42, 0.92, 0.28, size=9.5, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "←", 3.35, 3.68, 0.62, 0.5, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "路由 / 鉴权", 8.18, 2.22, 0.86, 0.28, size=9.5, color=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 8.3, 2.48, 0.62, 0.5, size=25, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "状态 / 结果", 8.18, 3.42, 0.86, 0.28, size=9.5, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "←", 8.3, 3.68, 0.62, 0.5, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 0.98, 5.62, 11.36, 0.62, fill=LIGHT_BLUE, line=RGBColor(0xB9, 0xD4, 0xE8))
    add_text(
        slide,
        "一人一账户 · 一机一设备身份 · 数据留在电脑端 · 高风险操作二次确认 · 支持随时撤销",
        1.12,
        5.69,
        11.08,
        0.44,
        size=14.5,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    return slide


def add_value_lane(
    slide,
    number: str,
    title: str,
    description: str,
    outcomes: str,
    y: float,
    *,
    accent: RGBColor,
    fill: RGBColor,
) -> None:
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(3.7),
        Inches(y + 0.24),
        Inches(0.5),
        Inches(0.5),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.color.rgb = accent
    add_text(slide, number, 3.7, y + 0.24, 0.5, 0.5, size=12, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, 4.4, y + 0.12, 1.75, 0.34, size=16, color=DARK_BLUE, bold=True)
    add_text(slide, description, 4.4, y + 0.5, 3.72, 0.48, size=11.5, color=TEXT)

    arrow = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON,
        Inches(8.28),
        Inches(y + 0.35),
        Inches(0.46),
        Inches(0.3),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = accent
    arrow.line.fill.background()

    add_round_rect(slide, 8.96, y + 0.12, 3.62, 0.86, fill=fill, line=fill)
    add_text(slide, outcomes, 9.16, y + 0.2, 3.22, 0.68, size=12.5, color=accent, bold=True, align=PP_ALIGN.CENTER)


def build_application_value_slide(prs: Presentation, layout=None):
    slide = prs.slides.add_slide(layout or prs.slide_layouts[3])
    clear_body(slide)
    set_title(slide, "应用价值：从一次性响应走向持续协作基础设施")
    add_subtitle(slide, "以可信长期记忆连接个人、组织与国产终端，让交互经验持续沉淀、复用并受控流转")

    add_text(slide, "能力底座", 0.72, 1.34, 2.58, 0.32, size=13, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "应用场景", 3.7, 1.34, 4.48, 0.32, size=13, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "价值结果", 8.96, 1.34, 3.62, 0.32, size=13, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 0.72, 1.72, 2.58, 3.78, fill=DARK_BLUE, line=DARK_BLUE, line_width=0.8)
    add_text(slide, "OSagent", 1.08, 2.0, 1.86, 0.48, size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "可信长期记忆底座", 1.0, 2.45, 2.02, 0.38, size=13.5, color=RGBColor(0xD9, 0xEC, 0xF8), align=PP_ALIGN.CENTER)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.14),
        Inches(3.03),
        Inches(1.74),
        Inches(0.012),
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(0x74, 0xA9, 0xCF)
    divider.line.fill.background()
    for index, item in enumerate(("形成与更新", "检索与压缩", "遗忘与审计", "权限与插件治理")):
        y = 3.25 + index * 0.47
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.02), Inches(y + 0.12), Inches(0.1), Inches(0.1))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(0x79, 0xC6, 0xEC)
        dot.line.fill.background()
        add_text(slide, item, 1.26, y, 1.62, 0.34, size=11.5, color=WHITE)

    for y in (1.72, 3.02, 4.32):
        lane = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(3.52),
            Inches(y),
            Inches(9.06),
            Inches(1.18),
        )
        lane.fill.solid()
        lane.fill.fore_color.rgb = WHITE
        lane.line.color.rgb = RGBColor(0xD7, 0xE2, 0xEA)
        lane.line.width = Pt(0.9)

    add_value_lane(
        slide,
        "01",
        "个人长期协作",
        "跨会话记住偏好、任务状态与执行依据",
        "减少重复说明\n复杂任务连续推进",
        1.72,
        accent=BLUE,
        fill=LIGHT_BLUE,
    )
    add_value_lane(
        slide,
        "02",
        "组织知识沉淀",
        "将 SOP、专家经验与项目轨迹转化为可追溯资产",
        "经验可复用\n过程可追溯 · 交接更稳定",
        3.02,
        accent=GREEN,
        fill=LIGHT_GREEN,
    )
    add_value_lane(
        slide,
        "03",
        "国产端侧智能",
        "在银河麒麟本地连接模型、插件与系统 AI 能力",
        "数据留本地\n权限可控 · 行业能力可扩展",
        4.32,
        accent=RED,
        fill=LIGHT_RED,
    )

    add_round_rect(slide, 0.72, 5.78, 11.86, 0.5, fill=LIGHT_BLUE, line=RGBColor(0xB9, 0xD4, 0xE8))
    add_text(slide, "产品化路径", 0.98, 5.85, 1.18, 0.34, size=12, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "个人本地助理    ·    企业私有部署    ·    行业插件平台",
        2.28,
        5.83,
        9.96,
        0.36,
        size=14,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    return slide


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the OSagent defense PPT visual supplement.")
    parser.add_argument("--source", type=Path, required=True, help="Existing PowerPoint deck used as the theme source.")
    parser.add_argument("--output", type=Path, required=True, help="Non-destructive output PowerPoint path.")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--official",
        action="store_true",
        help="Insert a background slide, replace the existing results slide, and append the outlook slide.",
    )
    mode.add_argument(
        "--application-value",
        action="store_true",
        help="Keep every source slide unchanged and append one application-value slide.",
    )
    args = parser.parse_args()

    prs = Presentation(args.source)
    if args.application_value:
        source_slide_count = len(prs.slides)
        value_slide = build_application_value_slide(prs, layout=prs.slides[-1].slide_layout)
        add_official_frame(value_slide, source_slide_count + 1)
        rendered = str(source_slide_count + 1)
    elif args.official:
        results_slide = next(
            (slide for slide in prs.slides if "调优效果" in slide_text(slide)),
            None,
        )
        if results_slide is None:
            raise ValueError("The official deck does not contain a results slide.")
        build_official_results_slide(results_slide)

        engineering_slide = next(
            (slide for slide in prs.slides if "国产平台与工程闭环验证" in slide_text(slide)),
            None,
        )
        if engineering_slide is not None:
            replace_text_preserving_runs(engineering_slide, "291", "346")
            overlay_engineering_regression_card(engineering_slide)

        background_slide = prs.slides.add_slide(results_slide.slide_layout)
        build_background_slide(background_slide)
        move_slide(prs, background_slide, 0)
        add_official_frame(background_slide, 1)
        outlook_slide = build_official_mobile_outlook_slide(prs, layout=results_slide.slide_layout)
        add_official_frame(outlook_slide, len(prs.slides))
        rendered = "1,6,8"
    else:
        if len(prs.slides) < 9:
            raise ValueError("The source deck must contain at least nine slides.")
        build_background_slide(prs.slides[0])
        build_results_slide(prs.slides[8])
        build_outlook_slide(prs)
        rendered = "1,9,10"

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(f"saved={args.output.resolve()}")
    print(f"render_slides={rendered}")


if __name__ == "__main__":
    main()
